"""Render and export. These run LibreOffice, so they are the slow ones."""

import shutil

import pytest
from docx import Document

from config import ClientConfig
from fields import CLIENT_FIELDS
from render import RenderError, render

pytestmark = pytest.mark.skipif(
    not (shutil.which("soffice") or shutil.which("libreoffice")),
    reason="LibreOffice is needed to export a PDF",
)


def make_config(values, tmp_path, slug="acme"):
    return ClientConfig(slug=slug, path=tmp_path / f"{slug}.yaml", values=values)


def test_round_trip_produces_a_pdf_with_every_tag(
    client_values, cmo_values, tmp_path, template_path
):
    result = render(
        make_config(client_values, tmp_path),
        cmo_values,
        template_path=template_path,
        build_dir=tmp_path / "build",
        allow_scaffold=True,
    )
    assert result.pdf_path.is_file()
    assert result.page_count > 1
    assert set(result.tags_found) == {f.tag_name for f in CLIENT_FIELDS}


def test_client_values_reach_the_document(client_values, cmo_values, tmp_path, template_path):
    from pdfminer.high_level import extract_text

    result = render(
        make_config(client_values, tmp_path), cmo_values,
        template_path=template_path, build_dir=tmp_path / "build", allow_scaffold=True,
    )
    text = extract_text(str(result.pdf_path))
    normalised = " ".join(text.split())

    assert "Acme Sandbox Brands, LLC" in normalised
    assert "January 5, 2026" in normalised
    assert "$8,500 per month" in normalised
    assert "Weekly reporting" in normalised          # Schedule B loop ran


def test_cmo_block_is_prefilled(client_values, cmo_values, tmp_path, template_path):
    from pdfminer.high_level import extract_text

    result = render(
        make_config(client_values, tmp_path), cmo_values,
        template_path=template_path, build_dir=tmp_path / "build", allow_scaffold=True,
    )
    normalised = " ".join(extract_text(str(result.pdf_path)).split())
    assert "/s/ William Fikhman" in normalised
    assert "Title: CEO" in normalised


def test_no_jinja_syntax_survives_into_the_document(
    client_values, cmo_values, tmp_path, template_path
):
    result = render(
        make_config(client_values, tmp_path), cmo_values,
        template_path=template_path, build_dir=tmp_path / "build", allow_scaffold=True,
    )
    doc = Document(str(result.docx_path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "{%" not in text
    # The only {{...}} left standing must be SignNow tags.
    for fragment in text.split("{{")[1:]:
        name = fragment.split("}}")[0]
        assert name in {f.tag_name for f in CLIENT_FIELDS}, name


def test_template_carries_no_pending_marker(template_path):
    """The real agreement text is in; a PENDING marker would mean it regressed."""
    from docx import Document as ReadDocument
    from render import SCAFFOLD_MARKER

    doc = ReadDocument(str(template_path))
    assert SCAFFOLD_MARKER not in "\n".join(p.text for p in doc.paragraphs)


def test_executed_aminomega_agreement_is_reproduced(cmo_values, tmp_path, template_path):
    """Rendering the transcribed Aminomega config must give back the executed text."""
    import yaml
    from pdfminer.high_level import extract_text
    from tests.conftest import REPO_ROOT

    values = yaml.safe_load((REPO_ROOT / "clients" / "aminomega.yaml").read_text())
    values["signatory_email"] = "signer@example.com"   # not in the executed PDF
    result = render(
        make_config(values, tmp_path, slug="aminomega"), cmo_values,
        template_path=template_path, build_dir=tmp_path / "build",
    )
    text = " ".join(extract_text(str(result.pdf_path)).split())

    assert result.page_count == 8
    assert ("and Aminomega, LLC a Pennsylvania LLC with address of 794 Sunrise Blvd., "
            "Mt. Bethel, PA 18343 (“Company”)") in text
    assert ("Company will pay CMO a monthly service fee of $2,300 plus a commission "
            "equal to 10% of monthly Gross Amazon Sales.") in text
    assert "totaling $13,800, shall be paid in advance" in text
    assert "continues for an initial term of 6 months" in text
    assert "6. Bi-Weekly status meeting" in text          # Schedule B numbering
    assert "Name: William Fikhman Title: CEO" in text     # CMO block pre-filled


def test_entity_phrasing_follows_state_and_type(tmp_path):
    """"a Pennsylvania limited liability company" vs "an Illinois corporation"."""
    from conftest import VALID_CLIENT, VALID_CMO

    llc = dict(VALID_CLIENT, entity_type="LLC", state_of_incorporation="Pennsylvania")
    ctx = make_config(llc, tmp_path).context(VALID_CMO)
    assert (ctx["entity_article"], ctx["entity_type_long"]) == ("a", "limited liability company")

    corp = dict(VALID_CLIENT, entity_type="Corporation", state_of_incorporation="Illinois")
    ctx = make_config(corp, tmp_path).context(VALID_CMO)
    assert (ctx["entity_article"], ctx["entity_type_long"]) == ("an", "corporation")


def test_todo_in_cmo_config_blocks_the_build(
    client_values, cmo_values, tmp_path, template_path
):
    cmo_values["address"] = "TODO: CMO principal place of business"
    with pytest.raises(RenderError, match="TODO"):
        render(
            make_config(client_values, tmp_path), cmo_values,
            template_path=template_path, build_dir=tmp_path / "build", allow_scaffold=True,
        )


def test_an_unfilled_placeholder_stops_the_build(cmo_values, tmp_path):
    """A template variable with nothing to fill it must never reach a PDF."""
    from docx import Document as NewDocument

    template = tmp_path / "leaky.docx"
    doc = NewDocument()
    doc.add_paragraph("Client: {{ company_legal_name }}")
    doc.add_paragraph("Unmapped: {{ variable_nobody_supplies }}")
    doc.save(str(template))

    config = make_config({"company_legal_name": "Acme"}, tmp_path)
    with pytest.raises(RenderError, match="nothing supplies"):
        render(config, cmo_values, template_path=template, build_dir=tmp_path / "build")


def test_missing_template_is_reported(client_values, cmo_values, tmp_path):
    with pytest.raises(RenderError, match="template not found"):
        render(
            make_config(client_values, tmp_path), cmo_values,
            template_path=tmp_path / "nope.docx", build_dir=tmp_path / "build",
        )


def _tiny_png(path):
    """A valid 2x2 white PNG written with the standard library."""
    import struct, zlib

    def chunk(tag, data):
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    raw = b"".join(b"\x00" + b"\xff\xff\xff" * 2 for _ in range(2))
    png = (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
           + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))
    path.write_bytes(png)


def test_uspto_screenshot_is_placed_when_present(client_values, cmo_values, tmp_path, template_path):
    """William's rule: the trademark exhibit is the USPTO screenshot when captured."""
    from docx import Document as ReadDocument

    build = tmp_path / "build"
    shot = build / "acme" / "trademark-uspto.png"
    shot.parent.mkdir(parents=True)
    _tiny_png(shot)

    result = render(make_config(client_values, tmp_path), cmo_values,
                    template_path=template_path, build_dir=build, allow_scaffold=True)
    assert result.trademark_screenshot == shot

    doc = ReadDocument(str(result.docx_path))
    assert doc.inline_shapes, "the screenshot should be embedded as an inline image"
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "TBD" not in text


def test_exhibit_reads_tbd_without_a_screenshot(client_values, cmo_values, tmp_path, template_path):
    result = render(make_config(client_values, tmp_path), cmo_values,
                    template_path=template_path, build_dir=tmp_path / "build")
    assert result.trademark_screenshot is None
    from pdfminer.high_level import extract_text
    assert "TBD" in extract_text(str(result.pdf_path))
