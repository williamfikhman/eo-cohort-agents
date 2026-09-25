"""Fill the Word template for one client and export a PDF.

Two things are verified before the PDF is handed on for upload: that no
placeholder survived the fill, and that every SignNow tag is still findable in the
PDF's text layer. A tag LibreOffice split across runs would upload cleanly and
silently produce a document with no signature field, which is the failure mode
worth spending a check on.
"""

from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from docxtpl import DocxTemplate, InlineImage
from docx.shared import Inches
from jinja2 import Environment, StrictUndefined
from jinja2.exceptions import UndefinedError

from config import ClientConfig
from fields import CLIENT_FIELDS, tag_context

REPO_ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_PATH = REPO_ROOT / "templates" / "amazon_services_agreement.docx"
BUILD_DIR = REPO_ROOT / "build"

#: Marker left in the scaffold template where real clause text belongs.
SCAFFOLD_MARKER = "CLAUSE TEXT PENDING"
#: Marker for CMO config values that were never filled in.
TODO_MARKER = "TODO:"

_PLACEHOLDER = re.compile(r"\{\{.*?\}\}|\{%.*?%\}", re.S)
_SOFFICE_CANDIDATES = ("soffice", "libreoffice")


class RenderError(Exception):
    """Raised when an agreement cannot be built, or was built wrong."""


@dataclass(frozen=True)
class RenderResult:
    slug: str
    docx_path: Path
    pdf_path: Path
    page_count: int
    tags_found: tuple[str, ...]
    trademark_screenshot: Path | None = None


def _docx_text(path: Path) -> str:
    """All text in the document, including tables, headers and footers."""
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(p.text for p in cell.paragraphs)
    for section in doc.sections:
        for container in (section.header, section.footer):
            parts.extend(p.text for p in container.paragraphs)
    return "\n".join(parts)


def _find_soffice() -> str:
    for name in _SOFFICE_CANDIDATES:
        found = shutil.which(name)
        if found:
            return found
    raise RenderError(
        "LibreOffice not found. Install it and make sure `soffice` is on PATH; "
        "the PDF export depends on it."
    )


def _pdf_text(path: Path) -> str:
    from pdfminer.high_level import extract_text

    try:
        return extract_text(str(path)) or ""
    except Exception as exc:  # pragma: no cover - depends on local pdf tooling
        raise RenderError(f"could not read text back out of {path}: {exc}") from exc


def _pdf_page_count(path: Path) -> int:
    from pdfminer.high_level import extract_pages

    return sum(1 for _ in extract_pages(str(path)))


def export_pdf(docx_path: Path, out_dir: Path) -> Path:
    """Convert to PDF with LibreOffice headless."""
    out_dir.mkdir(parents=True, exist_ok=True)
    soffice = _find_soffice()

    # LibreOffice needs a private profile or it silently no-ops when another
    # instance has the default profile locked.
    with tempfile.TemporaryDirectory(prefix="agreement-lo-") as profile:
        cmd = [
            soffice,
            "--headless",
            "--norestore",
            "--invisible",
            f"-env:UserInstallation=file://{profile}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(out_dir),
            str(docx_path),
        ]
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=300)

    pdf_path = out_dir / (docx_path.stem + ".pdf")
    if proc.returncode != 0 or not pdf_path.is_file():
        combined = f"{proc.stdout} {proc.stderr}"
        hint = ""
        if "source file could not be loaded" in combined:
            # LibreOffice reports a missing import filter the same way it reports
            # a corrupt file, and core installs without Writer more often than
            # you would expect.
            hint = (
                "\n\n  LibreOffice says it could not load the file. If the .docx "
                "opens fine\n  elsewhere, the Writer component is probably missing "
                "-- libreoffice-core\n  alone cannot read .docx. Install it:\n"
                "      sudo apt-get install libreoffice-writer   # Debian/Ubuntu\n"
                "      brew install --cask libreoffice           # macOS"
            )
        raise RenderError(
            "LibreOffice failed to export a PDF.\n"
            f"  exit code: {proc.returncode}\n"
            f"  stdout: {proc.stdout.strip()}\n"
            f"  stderr: {proc.stderr.strip()}{hint}"
        )
    return pdf_path


def render(
    client: ClientConfig,
    cmo: dict[str, Any],
    template_path: Path | None = None,
    build_dir: Path | None = None,
    allow_scaffold: bool = False,
    signnow_tags: bool = True,
) -> RenderResult:
    """Fill the template, export a PDF, and verify both.

    ``signnow_tags=False`` renders the tag anchors as nothing at all, for a
    clean PDF that will be uploaded and have its fields placed by hand.
    """
    template = template_path or TEMPLATE_PATH
    if not template.is_file():
        raise RenderError(f"template not found: {template}")

    out_dir = (build_dir or BUILD_DIR) / client.slug
    out_dir.mkdir(parents=True, exist_ok=True)

    context = client.context(cmo)
    context.update(tag_context() if signnow_tags else {k: "" for k in tag_context()})

    todo_values = sorted(
        f"{k}={v!r}" for k, v in context.items()
        if isinstance(v, str) and TODO_MARKER in v
    )
    if todo_values:
        listed = "\n".join(f"  - {v}" for v in todo_values)
        raise RenderError(
            "config still contains TODO placeholders:\n"
            f"{listed}\n\nFill these in before building an agreement."
        )

    doc = DocxTemplate(str(template))

    # Schedule C exhibit: the USPTO trademark screenshot captured by
    # tools/uspto_trademark.py, when it exists; otherwise the config's text
    # (usually "TBD"). William's rule is that every agreement carries the
    # screenshot, so its absence is reported by the CLI, not hidden.
    shot = out_dir / "trademark-uspto.png"
    if shot.is_file():
        context["trademark_exhibit"] = InlineImage(doc, str(shot), width=Inches(6))

    # StrictUndefined, not Jinja's default: an undefined variable must be a loud
    # failure, never a silently empty line in a signed contract.
    try:
        doc.render(context, Environment(undefined=StrictUndefined))
    except UndefinedError as exc:
        raise RenderError(
            f"the template refers to a variable that nothing supplies: {exc}\n\n"
            "Either add it to the client config or remove it from "
            f"{template.name}. Nothing was exported or uploaded."
        ) from exc

    docx_path = out_dir / f"{client.slug}-amazon-services-agreement.docx"
    doc.save(str(docx_path))

    text = _docx_text(docx_path)

    leftovers = sorted(set(_PLACEHOLDER.findall(text)) - {f.literal for f in CLIENT_FIELDS})
    if leftovers:
        listed = "\n".join(f"  - {p}" for p in leftovers)
        raise RenderError(
            f"the rendered agreement still contains unfilled placeholders:\n{listed}\n\n"
            "Nothing was exported or uploaded."
        )

    if SCAFFOLD_MARKER in text and not allow_scaffold:
        raise RenderError(
            "the template is still the generated scaffold -- its clause bodies read\n"
            f"  '{SCAFFOLD_MARKER}...'\n\n"
            "Paste the real agreement text into templates/amazon_services_agreement.docx\n"
            "before sending this to a client. To build it anyway for a sandbox test, "
            "pass --allow-scaffold."
        )

    pdf_path = export_pdf(docx_path, out_dir)
    pdf_text = _pdf_text(pdf_path)

    missing = [f.tag_name for f in CLIENT_FIELDS if f.literal not in pdf_text]
    if missing and signnow_tags:
        raise RenderError(
            "these SignNow tags did not survive the PDF export intact: "
            f"{', '.join(missing)}\n\n"
            "SignNow finds fields by matching this text, so uploading now would "
            "produce a document with missing signature fields. This usually means "
            "the tag got split across formatting runs in Word -- retype it as a "
            "single unstyled run."
        )

    return RenderResult(
        slug=client.slug,
        docx_path=docx_path,
        pdf_path=pdf_path,
        page_count=_pdf_page_count(pdf_path),
        tags_found=tuple(f.tag_name for f in CLIENT_FIELDS) if signnow_tags else (),
        trademark_screenshot=shot if shot.is_file() else None,
    )
